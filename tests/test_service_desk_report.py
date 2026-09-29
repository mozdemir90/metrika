import io
import openpyxl
import pandas as pd
import pytest
from app.reports.service_desk_report import ServiceDeskReport


@pytest.fixture
def report_engine():
    return ServiceDeskReport()


def test_technician_name_normalization(report_engine):
    assert report_engine.normalize_technician_name("glpi \nAhmet Yılmaz") == "Ahmet Yılmaz"
    assert report_engine.normalize_technician_name("glpi \r\n  Mehmet Kaya  ") == "Mehmet Kaya"
    assert report_engine.normalize_technician_name("glpi\nBurak Şahin") == "Burak Şahin"
    assert report_engine.normalize_technician_name("Ayşe Demir") == "Ayşe Demir"
    assert report_engine.normalize_technician_name(None) == "Atanmamış"
    assert report_engine.normalize_technician_name("") == "Atanmamış"
    assert report_engine.normalize_technician_name("nan") == "Atanmamış"


def test_split_technician_names(report_engine):
    assert report_engine.split_technician_names("glpi \nBurak Çıkrık\nHasan Karakaşoğlu") == ["Burak Çıkrık", "Hasan Karakaşoğlu"]
    assert report_engine.split_technician_names("Burak Çıkrık \r\nMerve Esen Baygüneş") == ["Burak Çıkrık", "Merve Esen Baygüneş"]
    assert report_engine.split_technician_names("glpi \nTek Kişi") == ["Tek Kişi"]
    assert report_engine.split_technician_names(None) == ["Atanmamış"]
    assert report_engine.split_technician_names("") == ["Atanmamış"]


def test_duration_parsing(report_engine):
    # 1 day = 1440 min, 4 hours = 240 min, 30 min = 30 min -> 1710 min = 28.5 hours
    mins = report_engine.parse_duration_to_minutes("1 gün 4 saat 30 dakika")
    assert mins == 1710.0
    hours = report_engine.parse_duration_to_hours("1 gün 4 saat 30 dakika")
    assert hours == 28.5

    assert report_engine.parse_duration_to_minutes("45 dakika") == 45.0
    assert report_engine.parse_duration_to_minutes("2 saat 15 dk") == 135.0
    assert report_engine.parse_duration_to_minutes("0 dakika") == 0.0
    assert report_engine.parse_duration_to_minutes(None) == 0.0


def test_calculate_metrics_business_logic(report_engine):
    data = [
        {
            "ticket_id": 1,
            "title": "Adalet Bakanlığı - E-İmza",
            "institution": "Adalet Bakanlığı",
            "category": "E-İmza",
            "root_cause": "Hata",
            "priority": "Yüksek",
            "status": "Çözüldü",
            "technician_raw": "glpi \nAhmet Yılmaz",
            "created_at": "2026-09-01 10:00:00",
            "resolved_at": "2026-09-01 18:00:00",
            "first_response_time_str": "20 dakika",
            "total_resolution_time_str": "8 saat",
            "external_wait_time_str": "2 saat",
            "is_fcr": True,
            "is_reopened": False,
        },
        {
            "ticket_id": 2,
            "title": "Milli Eğitim - Form Düzenleme",
            "institution": "Milli Eğitim Bakanlığı",
            "category": "Şablon",
            "root_cause": "Destek",
            "priority": "Orta",
            "status": "Çözüldü",
            "technician_raw": "Ayşe Demir",
            "created_at": "2026-09-02 09:00:00",
            "resolved_at": "2026-09-03 15:00:00",
            "first_response_time_str": "50 dakika",
            "total_resolution_time_str": "30 saat",
            "external_wait_time_str": "0 dakika",
            "is_fcr": False,
            "is_reopened": True,
        },
    ]

    df = pd.DataFrame(data)
    metrics = report_engine.calculate_metrics(df)

    summary = metrics["summary"]
    assert summary["total_tickets"] == 2
    assert summary["resolved_tickets"] == 2
    assert summary["bug_count"] == 1
    assert summary["support_count"] == 1

    # MTTA SLA: Ticket 1 High (20 min <= 30 min -> PASS), Ticket 2 Med (50 min <= 60 min -> PASS)
    assert summary["mtta_sla_rate"] == 100.0

    # Net resolution for Ticket 1: 8h - 2h = 6h (<= 24h -> PASS)
    # Net resolution for Ticket 2: 30h - 0h = 30h (<= 48h -> PASS)
    assert summary["mttr_sla_rate"] == 100.0

    # FCR: 1 out of 2 = 50%
    assert summary["fcr_rate"] == 50.0

    # Reopen: 1 out of 2 = 50%
    assert summary["reopen_rate"] == 50.0

    # Technicians
    tech_names = [t["technician"] for t in metrics["technicians"]]
    assert "Ahmet Yılmaz" in tech_names
    assert "Ayşe Demir" in tech_names


def test_excel_rendering(report_engine):
    data = [
        {
            "ticket_id": 101,
            "title": "Test Ticket",
            "institution": "Sağlık Bakanlığı",
            "category": "Test",
            "root_cause": "Hata",
            "priority": "Yüksek",
            "status": "Çözüldü",
            "technician_raw": "glpi \nCanan Öztürk",
            "created_at": "2026-09-01 10:00:00",
            "resolved_at": "2026-09-01 14:00:00",
            "first_response_time_str": "15 dakika",
            "total_resolution_time_str": "4 saat",
            "external_wait_time_str": "1 saat",
            "is_fcr": True,
            "is_reopened": False,
        }
    ]
    df = pd.DataFrame(data)
    metrics = report_engine.calculate_metrics(df)

    excel_buf = report_engine.render_excel(metrics, df)
    assert isinstance(excel_buf, io.BytesIO)
    assert excel_buf.getvalue() != b""

    # Load back with openpyxl to verify workbook structure
    wb = openpyxl.load_workbook(excel_buf)
    sheet_names = wb.sheetnames
    assert len(sheet_names) == 6
    assert "1. Konsolide Yönetici Özeti" in sheet_names
    assert "2. Servis Masası (L1)" in sheet_names
    assert "3. Sistem Ekibi (L2)" in sheet_names
    assert "4. Kurum & Talep Analizi" in sheet_names
    assert "5. Detaylı Çağrı Listesi" in sheet_names
    assert "6. OpenProject & Kurum Analizi" in sheet_names

    ws_summary = wb["1. Konsolide Yönetici Özeti"]
    assert "METRİKA" in ws_summary["B2"].value


def test_glpi_real_report_hk_file(report_engine):
    import os
    file_path = "GLPI RAPOR HK.xlsx"
    if not os.path.exists(file_path):
        pytest.skip("GLPI RAPOR HK.xlsx not found")

    df = pd.read_excel(file_path)
    metrics = report_engine.calculate_metrics(df, segment="all")

    summary = metrics["summary"]
    assert summary["total_tickets"] == 139
    assert summary["resolved_tickets"] == 54
    assert summary["reopen_count"] == 4
    assert summary["fcr_count"] == 15

    # Team segmentation verification
    tb = metrics["team_breakdown"]
    assert tb["service_desk_count"] == 139
    assert tb["system_ops_count"] == 43

    # Segment specific calculation test
    metrics_sd = report_engine.calculate_metrics(df, segment="service_desk")
    assert metrics_sd["summary"]["total_tickets"] == 139

    metrics_sys = report_engine.calculate_metrics(df, segment="system_ops")
    assert metrics_sys["summary"]["total_tickets"] == 43

    # Verify technician names are present and correctly parsed
    techs = [t["technician"] for t in metrics["technicians"]]
    assert "İsmail Demir" in techs
    assert "Orhan Erdoğan" in techs
    assert "Ahmet Arif AKBAŞ" in techs

    # Check top institution
    assert len(metrics["top_institutions"]) > 0

    # Test excel render (6 sheets)
    buf = report_engine.render_excel(metrics, df)
    wb = openpyxl.load_workbook(buf)
    assert len(wb.sheetnames) == 6
    assert "2. Servis Masası (L1)" in wb.sheetnames
    assert "3. Sistem Ekibi (L2)" in wb.sheetnames
    assert "6. OpenProject & Kurum Analizi" in wb.sheetnames
    buf = report_engine.render_excel(metrics, df)
    assert buf.getvalue() != b""


def test_escalation_flow_metrics(report_engine):
    """Verifies that L1 directly resolved tickets (FCR) and L2/L3 escalations are strictly partitioned."""
    rows = []
    # 32 L2 tickets (Sistem / Veritabanı / Deploy)
    for i in range(29):
        rows.append({"ticket_id": f"BN-VT-{i}", "op_task_type": "BN-Veritabanı İşlemi Görevi", "status": "Çözüldü", "priority": "Orta"})
    for i in range(3):
        rows.append({"ticket_id": f"BN-DEP-{i}", "op_task_type": "BN-Deploy (UGD)", "status": "Çözüldü", "priority": "Orta"})

    # 50 L3 tickets (Belgenet Hata Kurum & Öneri)
    for i in range(42):
        rows.append({"ticket_id": f"BNSM-HATA-{i}", "op_task_type": "BN-Belgenet Hata Kurum", "status": "İşleniyor", "priority": "Yüksek"})
    for i in range(8):
        rows.append({"ticket_id": f"BNSM-ONR-{i}", "op_task_type": "BN-Belgenet Öneri", "status": "Bekliyor", "priority": "Düşük"})

    # 22 L1 Direct tickets (no OP task or basket): 15 resolved, 7 open
    for i in range(12):
        rows.append({"ticket_id": f"BNSM-DIR-RES-{i}", "op_task_type": "", "op_basket": "", "status": "Çözülmüş", "priority": "Orta"})
    for i in range(3):
        rows.append({"ticket_id": f"BNSM-DIR-CLS-{i}", "op_task_type": "", "op_basket": "", "status": "Kapalı", "priority": "Orta"})
    for i in range(4):
        rows.append({"ticket_id": f"BNSM-DIR-WT-{i}", "op_task_type": "", "op_basket": "", "status": "Bekliyor", "priority": "Düşük"})
    for i in range(3):
        rows.append({"ticket_id": f"BNSM-DIR-IP-{i}", "op_task_type": "", "op_basket": "", "status": "İşleniyor (atanmış)", "priority": "Orta"})

    df = pd.DataFrame(rows)
    assert len(df) == 104

    metrics = report_engine.calculate_metrics(df)
    esc = metrics["escalation_summary"]

    assert esc["total_incoming"] == 104
    assert esc["l1_direct"]["count"] == 15
    assert esc["l1_direct"]["resolved"] == 15
    assert esc["l1_direct"]["total_unforwarded"] == 22
    assert esc["l1_direct"]["open"] == 7
    assert esc["l1_direct"]["rate"] == 14.4
    assert esc["l1_direct"]["resolution_rate"] == 68.2

    assert esc["l2_transfers"]["count"] == 32
    assert esc["l2_transfers"]["rate"] == 30.8

    # L3 Yazılım / Hata (Öneriler ayrılmış: 42 adet)
    assert esc["l3_transfers"]["count"] == 42
    assert esc["l3_transfers"]["rate"] == 40.4
    assert "BN-Belgenet Hata Kurum" in esc["l3_transfers"]["task_types"]
    assert "Belirtilmemiş" not in esc["l3_transfers"]["task_types"]

    # İş Geliştirme (Öneri olanlar: 8 adet)
    assert "business_dev_transfers" in esc
    assert esc["business_dev_transfers"]["count"] == 8
    assert esc["business_dev_transfers"]["rate"] == 7.7
    assert "BN-Belgenet Öneri" in esc["business_dev_transfers"]["task_types"]


