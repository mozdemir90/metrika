import io
import os
import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


@pytest.fixture(autouse=True)
def preserve_uploaded_dataset():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target_path = os.path.join(base_dir, "last_uploaded_dataset.xlsx")
    existing_bytes = None
    if os.path.exists(target_path):
        with open(target_path, "rb") as f:
            existing_bytes = f.read()
    yield
    if existing_bytes is not None:
        with open(target_path, "wb") as f:
            f.write(existing_bytes)
    elif os.path.exists(target_path):
        try:
            os.remove(target_path)
        except OSError:
            pass


def test_dashboard_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    assert "METRİKA" in response.text


def test_health_check_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "uptime_seconds" in data
    assert "data_sources" in data


def test_report_types_endpoint():
    response = client.get("/api/v1/reports/types")
    assert response.status_code == 200
    data = response.json()
    assert "reports" in data
    assert "connectors" in data
    assert any(r["key"] == "service_desk" for r in data["reports"])


def test_get_metrics_mock_endpoint():
    response = client.get("/api/v1/reports/service_desk/metrics?days=7&source=mock")
    assert response.status_code == 200
    data = response.json()
    assert "summary" in data
    assert "priorities" in data
    assert "technicians" in data
    assert "top_institutions" in data
    assert "executive_insights" in data

    summary = data["summary"]
    assert summary["total_tickets"] > 0
    assert "overall_sla_rate" in summary
    assert "fcr_rate" in summary
    assert "reopen_rate" in summary


def test_export_excel_endpoint():
    response = client.get("/api/v1/reports/service_desk/export-excel?days=7&source=mock")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Content-Disposition" in response.headers
    assert "METRIKA_Belgenet" in response.headers["Content-Disposition"]

    # Verify content is valid Excel with 5 corporate sheets
    excel_buf = io.BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_buf)
    assert "1. Konsolide Yönetici Özeti" in wb.sheetnames
    assert "2. Servis Masası (L1)" in wb.sheetnames
    assert "3. Sistem Ekibi (L2)" in wb.sheetnames
    assert "4. Kurum & Talep Analizi" in wb.sheetnames
    assert "5. Detaylı Çağrı Listesi" in wb.sheetnames


def test_segment_metrics_and_export():
    # Test segment parameter in metrics endpoint
    res_sd = client.get("/api/v1/reports/service_desk/metrics?source=mock&segment=service_desk")
    assert res_sd.status_code == 200
    data_sd = res_sd.json()
    assert data_sd["active_segment"] == "service_desk"
    assert "summary" in data_sd

    res_sys = client.get("/api/v1/reports/service_desk/metrics?source=mock&segment=system_ops")
    assert res_sys.status_code == 200
    data_sys = res_sys.json()
    assert data_sys["active_segment"] == "system_ops"

    # Test segment export
    res_export = client.get("/api/v1/reports/service_desk/export-excel?source=mock&segment=service_desk")
    assert res_export.status_code == 200
    wb_sd = openpyxl.load_workbook(io.BytesIO(res_export.content))
    assert "1. Konsolide Yönetici Özeti" in wb_sd.sheetnames


def test_upload_excel_dataset():
    # Create sample in-memory excel
    df = pd.DataFrame([
        {
            "Çağrı No": 501,
            "Başlık": "E-İmza Sorunu",
            "Kurum Adı": "Adalet Bakanlığı",
            "Kategori": "E-İmza",
            "Öncelik": "Yüksek",
            "Durum": "Çözüldü",
            "Teknisyen": "glpi \nAhmet Yılmaz",
            "İlk Yanıt Süresi": "15 dakika",
            "Toplam Çözüm Süresi": "6 saat",
            "Dış Bekleme": "1 saat",
            "FCR": "Evet",
            "Yeniden Açıldı": "Hayır",
        },
        {
            "Çağrı No": 502,
            "Başlık": "EBYS İstemci Kurulumu",
            "Kurum Adı": "Sağlık Bakanlığı",
            "Kategori": "Kurulum",
            "Öncelik": "Orta",
            "Durum": "Çözüldü",
            "Teknisyen": "Ayşe Demir",
            "İlk Yanıt Süresi": "30 dakika",
            "Toplam Çözüm Süresi": "12 saat",
            "Dış Bekleme": "2 saat",
            "FCR": "Hayır",
            "Yeniden Açıldı": "Evet",
        }
    ])
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)

    response = client.post(
        "/api/v1/reports/service_desk/upload",
        files={"file": ("test_glpi_export.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    res_data = response.json()
    assert "metrics" in res_data
    summary = res_data["metrics"]["summary"]
    assert summary["total_tickets"] == 2
    assert summary["fcr_count"] == 1
    assert summary["reopen_count"] == 1

    # Now verify querying with source=upload
    res_upload_metrics = client.get("/api/v1/reports/service_desk/metrics?source=upload")
    assert res_upload_metrics.status_code == 200
    assert res_upload_metrics.json()["summary"]["total_tickets"] == 2


def test_export_from_upload_direct():
    csv_content = (
        "ticket_id,title,institution,category,root_cause,priority,status,technician_raw,first_response_time_str,total_resolution_time_str,external_wait_time_str,is_fcr,is_reopened\n"
        "999,Sunucu Kesintisi,İçişleri Bakanlığı,Altyapı,Hata,Yüksek,Çözüldü,glpi \\nMehmet Kaya,20 dakika,4 saat,1 saat,True,False\n"
    )
    response = client.post(
        "/api/v1/reports/service_desk/export-from-upload",
        files={"file": ("raw.csv", csv_content.encode("utf-8"), "text/csv")}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "1. Konsolide Yönetici Özeti" in wb.sheetnames


def test_upload_dataset_date_filtering():
    """
    Verifies that when source=upload, filtering by days (7, 15, 30, 0)
    accurately subsets the data relative to the maximum date in the file.
    """
    df = pd.DataFrame([
        {
            "Çağrı No": 1001,
            "Başlık": "Bugünlük Çağrı",
            "Kurum Adı": "Kurum A",
            "Açılış tarihi": "22-09-2026 14:00",
            "Kategori": "Yazılım",
            "Öncelik": "Yüksek",
            "Durum": "Çözüldü",
            "Teknisyen": "Burak Çıkrık",
            "İlk Yanıt Süresi": "10 dakika",
            "Toplam Çözüm Süresi": "2 saat",
            "Dış Bekleme": "0 saat",
            "FCR": "Evet",
            "Yeniden Açıldı": "Hayır",
        },
        {
            "Çağrı No": 1002,
            "Başlık": "10 Gün Önceki Çağrı",
            "Kurum Adı": "Kurum B",
            "Açılış tarihi": "12-09-2026 10:00",
            "Kategori": "Donanım",
            "Öncelik": "Orta",
            "Durum": "Çözüldü",
            "Teknisyen": "Hasan Karakaşoğlu",
            "İlk Yanıt Süresi": "20 dakika",
            "Toplam Çözüm Süresi": "5 saat",
            "Dış Bekleme": "1 saat",
            "FCR": "Hayır",
            "Yeniden Açıldı": "Hayır",
        },
        {
            "Çağrı No": 1003,
            "Başlık": "25 Gün Önceki Çağrı",
            "Kurum Adı": "Kurum C",
            "Açılış tarihi": "28-08-2026 09:00",
            "Kategori": "E-İmza",
            "Öncelik": "Düşük",
            "Durum": "Çözüldü",
            "Teknisyen": "Merve Esen Baygüneş",
            "İlk Yanıt Süresi": "30 dakika",
            "Toplam Çözüm Süresi": "10 saat",
            "Dış Bekleme": "2 saat",
            "FCR": "Evet",
            "Yeniden Açıldı": "Evet",
        },
    ])
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)

    # Upload dataset
    up_res = client.post(
        "/api/v1/reports/service_desk/upload",
        files={"file": ("test_multi_date.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert up_res.status_code == 200

    # Test 7 days filter -> only record 1 (1001)
    res_7 = client.get("/api/v1/reports/service_desk/metrics?source=upload&days=7")
    assert res_7.status_code == 200
    assert res_7.json()["summary"]["total_tickets"] == 1

    # Test 15 days filter -> records 1 & 2 (1001, 1002)
    res_15 = client.get("/api/v1/reports/service_desk/metrics?source=upload&days=15")
    assert res_15.status_code == 200
    assert res_15.json()["summary"]["total_tickets"] == 2

    # Test 30 days filter -> all 3 records
    res_30 = client.get("/api/v1/reports/service_desk/metrics?source=upload&days=30")
    assert res_30.status_code == 200
    assert res_30.json()["summary"]["total_tickets"] == 3

    # Test 0 days filter (Tüm Veri) -> all 3 records
    res_0 = client.get("/api/v1/reports/service_desk/metrics?source=upload&days=0")
    assert res_0.status_code == 200
    assert res_0.json()["summary"]["total_tickets"] == 3

    # Test excel export with days=7 and source=upload
    res_exp = client.get("/api/v1/reports/service_desk/export-excel?source=upload&days=7")
    assert res_exp.status_code == 200
    assert "METRIKA_Belgenet_Servis_Masasi_Raporu_7Gun" in res_exp.headers["Content-Disposition"]
