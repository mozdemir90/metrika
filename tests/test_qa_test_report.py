import io
import openpyxl
import pandas as pd
import pytest
from app.reports.qa_test_report import QATestReport
from app.reports.report_factory import ReportFactory


def test_qa_report_factory_registration():
    report = ReportFactory.get_report("qa_testing")
    assert isinstance(report, QATestReport)

    types = ReportFactory.list_reports()
    assert any(t["key"] == "qa_testing" for t in types)


def test_qa_report_metrics_calculation():
    df = pd.DataFrame([
        {"Test Senaryosu": "TC-101", "Test Uzmanı": "Zeynep Kaya", "Durum": "Geçti", "Kritiklik": "Yüksek", "Efor": "2 saat"},
        {"Test Senaryosu": "TC-102", "Test Uzmanı": "Zeynep Kaya", "Durum": "Kaldı", "Kritiklik": "Kritik", "Efor": "4 saat"},
        {"Test Senaryosu": "TC-103", "Test Uzmanı": "Ali Can", "Durum": "Geçti", "Kritiklik": "Orta", "Efor": "1 saat"},
        {"Test Senaryosu": "TC-104", "Test Uzmanı": "Ali Can", "Durum": "Beklemede", "Kritiklik": "Düşük", "Efor": "1 saat"},
    ])

    qa_engine = QATestReport()
    metrics = qa_engine.calculate_metrics(df)

    summary = metrics["summary"]
    assert summary["total_tests"] == 4
    assert summary["passed_count"] == 2
    assert summary["failed_count"] == 1
    assert summary["in_progress_count"] == 1
    assert summary["pass_rate"] == 50.0

    testers = metrics["testers"]
    assert len(testers) == 2
    zeynep = next(t for t in testers if t["tester"] == "Zeynep Kaya")
    assert zeynep["total_executed"] == 2
    assert zeynep["passed_count"] == 1
    assert zeynep["failed_count"] == 1


def test_qa_report_excel_rendering():
    df = pd.DataFrame([
        {"Test Senaryosu": "TC-201", "Test Uzmanı": "Merve Yılmaz", "Durum": "Geçti", "Kritiklik": "Yüksek"},
        {"Test Senaryosu": "TC-202", "Test Uzmanı": "Merve Yılmaz", "Durum": "Geçti", "Kritiklik": "Orta"},
    ])
    qa_engine = QATestReport()
    metrics = qa_engine.calculate_metrics(df)
    excel_buf = qa_engine.render_excel(metrics, df)

    wb = openpyxl.load_workbook(excel_buf)
    assert "Test & QA Performans Özeti" in wb.sheetnames
    assert "Detaylı Test Koşumları" in wb.sheetnames
