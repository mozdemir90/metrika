import io
import os
import openpyxl
import pandas as pd
import pytest
from app.connectors.smart_column_mapper import SmartColumnMapper
from app.reports.service_desk_report import ServiceDeskReport

GLPI3_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "glpi(3).xlsx")


@pytest.fixture
def glpi3_df():
    assert os.path.exists(GLPI3_PATH), f"glpi(3).xlsx bulunamadı: {GLPI3_PATH}"
    return pd.read_excel(GLPI3_PATH)


@pytest.fixture
def report_engine():
    return ServiceDeskReport()


def test_smart_column_mapper_glpi3(glpi3_df):
    """SmartColumnMapper glpi(3).xlsx içindeki tüm standart ve OpenProject sütunlarını doğru tespit etmeli."""
    std_df, diagnostics = SmartColumnMapper.map_dataframe(glpi3_df)
    mapping = diagnostics["mapped_columns"]

    # Temel sütunlar
    assert "ticket_id" in mapping
    assert mapping["ticket_id"] in ("Kimlik", "Eklentiler - Semantic ID")
    assert "institution" in mapping
    assert mapping["institution"] == "İstekte bulunan - İstekte bulunan"
    assert "technician_raw" in mapping
    assert mapping["technician_raw"] == "Atananlar - Teknisyen"
    assert "technician_group" in mapping
    assert mapping["technician_group"] == "Atananlar - Teknisyen grubu"
    assert "total_resolution_time_str" in mapping
    assert mapping["total_resolution_time_str"] == "Çözümlenme süresi"

    # OpenProject sütunları
    assert "op_request_type" in mapping
    assert mapping["op_request_type"] == "Eklentiler - OpenProjecte gönderme - Talep Tipi"
    assert "op_task_type" in mapping
    assert mapping["op_task_type"] == "Eklentiler - OpenProjecte gönderme - Görev Tipi"
    assert "op_basket" in mapping
    assert mapping["op_basket"] == "Eklentiler - OpenProjecte gönderme - Bu destek kaydını OpenProjecte ilet"

    # Standardize edilmiş DF
    assert "op_request_type" in std_df.columns
    assert "op_task_type" in std_df.columns
    assert "op_basket" in std_df.columns
    assert len(std_df) == 1009


def test_service_desk_clean_data_glpi3(report_engine, glpi3_df):
    """ServiceDeskReport.standardize_dataframe 3-yönlü kök neden (Hata, Destek, Öneri) ve OP iletim bayrağını doğru işlemeli."""
    cleaned_df, _ = report_engine.standardize_dataframe(glpi3_df)

    # Kök neden 3 değer almalı
    root_causes = cleaned_df["root_cause"].value_counts()
    assert "Hata" in root_causes
    assert "Destek" in root_causes
    assert "Öneri" in root_causes
    assert root_causes["Hata"] == 94
    assert root_causes["Destek"] == 878
    assert root_causes["Öneri"] == 37
    assert len(cleaned_df) == 1009

    # OpenProject iletim sayısı
    op_forwarded = cleaned_df[cleaned_df["is_op_forwarded"] == True]
    assert len(op_forwarded) == 184


def test_service_desk_metrics_glpi3(report_engine, glpi3_df):
    """ServiceDeskReport.calculate_metrics kurum bazlı OpenProject metriklerini ve çapraz matrisi doğru hesaplamalı."""
    cleaned_df, _ = report_engine.standardize_dataframe(glpi3_df)
    metrics = report_engine.calculate_metrics(cleaned_df, segment="all")

    summary = metrics["summary"]
    assert summary["total_tickets"] == 1009
    assert summary["bug_count"] == 94
    assert summary["support_count"] == 878
    assert summary["suggestion_count"] == 37
    assert summary["op_forwarded_total"] == 184

    # Global OpenProject özeti
    assert "openproject_summary" in metrics
    op_summary = metrics["openproject_summary"]
    assert op_summary["total_forwarded"] == 184
    assert len(op_summary["baskets"]) > 0
    assert len(op_summary["task_types"]) > 0
    assert len(op_summary["demand_types"]) > 0

    # Sepet isimleri kontrolü
    basket_names = list(op_summary["baskets"].keys())
    assert any("Hata" in b for b in basket_names)
    assert any("İyileştirme" in b for b in basket_names)
    assert any("Sistem" in b for b in basket_names)

    # Kurum bazlı analizler
    assert "all_institutions" in metrics
    all_insts = metrics["all_institutions"]
    assert len(all_insts) > 50

    cb = next(i for i in all_insts if "Cumhurbaşkanlığı" in i["institution"])
    assert cb["total_count"] == 153
    assert cb["op_forwarded_count"] == 3
    assert cb["suggestion_count"] == 2
    assert len(cb["openproject_summary"]["cross_matrix"]) > 0
    assert len(cb["tickets_preview"]) > 0


def test_service_desk_excel_export_6_sheets(report_engine, glpi3_df):
    """ServiceDeskReport.render_excel 6 sekmeli ve OpenProject & Kurum Analizi içeren Excel üretmeli."""
    cleaned_df, _ = report_engine.standardize_dataframe(glpi3_df)
    metrics = report_engine.calculate_metrics(cleaned_df, segment="all")
    excel_buf = report_engine.render_excel(metrics, cleaned_df)

    wb = openpyxl.load_workbook(excel_buf)
    sheet_names = wb.sheetnames

    # 6 sekme mevcut olmalı
    assert len(sheet_names) == 6
    assert "1. Konsolide Yönetici Özeti" in sheet_names
    assert "2. Servis Masası (L1)" in sheet_names
    assert "3. Sistem Ekibi (L2)" in sheet_names
    assert "4. Kurum & Talep Analizi" in sheet_names
    assert "5. Detaylı Çağrı Listesi" in sheet_names
    assert "6. OpenProject & Kurum Analizi" in sheet_names

    # Sheet 4 kontrolü (Öneri ve OP sütunları - row 4 header)
    ws4 = wb["4. Kurum & Talep Analizi"]
    headers_ws4 = [ws4.cell(row=4, column=c).value for c in range(1, 9)]
    assert "Öneri (Ek Geliştirme)" in headers_ws4
    assert "OP İletilen" in headers_ws4

    # Sheet 5 kontrolü (OP sütunları - row 1 header, 20 sütun)
    ws5 = wb["5. Detaylı Çağrı Listesi"]
    headers_ws5 = [ws5.cell(row=1, column=c).value for c in range(1, 21)]
    assert "OP Talep Tipi" in headers_ws5
    assert "OP Görev Tipi" in headers_ws5
    assert "OP Hedef Sepet" in headers_ws5

    # Sheet 6 kontrolü (Başlık, KPI'lar ve Tablolar)
    ws6 = wb["6. OpenProject & Kurum Analizi"]
    title_val = ws6.cell(row=2, column=2).value
    assert "OPENPROJECT" in title_val

    # Table 1 başlıkları (row 10, cols 2-9)
    table1_headers = [ws6.cell(row=10, column=c).value for c in range(2, 10)]
    assert "Kurum Adı" in table1_headers
    assert "Toplam Çağrı" in table1_headers
    assert "OP'ye İletilen" in table1_headers
    assert "Öneri (Ek Geliştirme)" in table1_headers
